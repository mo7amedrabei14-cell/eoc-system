# -*- coding: utf-8 -*-
import sys, re
from pathlib import Path
from collections import Counter

try:
    from openpyxl import load_workbook
except ImportError:
    sys.exit("محتاج openpyxl")

SHEET = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(r"C:\Users\mo7am\OneDrive\Work\EOC System\تصنيفات مهام.xlsx")
NAMES = Path(sys.argv[2]) if len(sys.argv) > 2 else (Path.home() / "Downloads" / "names.txt")
OUT   = Path("missing.txt")

NAME_COL, VAL_COLS = 1, [2, 3, 4, 5, 6]
RULES_SHEET = "قواعد"
MIN_KEY_LEN = 3

DASH   = re.compile(r"[-\u2010-\u2015\u2043\u2212]")
DAYISH = re.compile(r"^(?:(?:ال|و)[\u0600-\u06FF]*|\d+)$")

def norm(s):
    s = str(s if s is not None else "")
    s = re.sub(r"[\u064B-\u0652\u0640\u0670]", "", s)
    s = re.sub(r"[أإآٱ]", "ا", s)
    s = re.sub(r"[ىئ]", "ي", s)
    s = re.sub(r"[ة]", "ه", s)
    s = re.sub(r"[\"“”«»'‘’]", " ", s)
    return re.sub(r"\s+", " ", DASH.sub("-", s)).strip()

def _strip_open_day(s):
    i = s.rfind("مفتوحه")
    if i == -1:
        return s
    tail = s[i + 6:].strip()
    toks = tail.split() if tail else []
    return s[:i].strip() if (toks and len(toks) <= 6 and all(DAYISH.match(t) for t in toks)) else s

def key_from(name):
    s = norm(name)
    for _ in range(3):
        s = re.sub(r"\s*مفتوحه\s*(اليوم.*)?$", "", s)
        s = re.sub(r"\s*\(\s*اليوم[^)]*\)\s*$", "", s)
        s = re.sub(r"\s*\d{4}\s*-\s*\d{1,2}\s*-\s*\d{1,2}\s*$", "", s).strip()
    s = _strip_open_day(s)
    parts = re.split(r"[اأإآ]\s*ستماره\s+مهمه", s)
    if len(parts) > 1:
        s = parts[-1].strip()
    return re.split(r"\s*-\s*", s, maxsplit=1)[0].strip()

def clean(v):
    return re.sub(r"\s+", " ", str(v)).strip() if v is not None else ""

def has_word(hay, needle):
    return f" {hay} ".find(f" {needle} ") != -1

if not SHEET.exists():
    sys.exit(f"مش لاقي ملف الشيت: {SHEET}")
if not NAMES.exists():
    sys.exit(f"مش لاقي ملف الأسماء: {NAMES}  (نزّله من المتصفح الأول)")

wb = load_workbook(SHEET, data_only=True)
keys = set()
for row in wb.worksheets[0].iter_rows(min_row=2, max_col=6, values_only=True):
    if not any(clean(row[c - 1]) for c in VAL_COLS):
        continue
    k = key_from(clean(row[NAME_COL - 1]))
    if len(k) >= MIN_KEY_LEN:
        keys.add(k)

rules = []
if RULES_SHEET in wb.sheetnames:
    for row in wb[RULES_SHEET].iter_rows(min_row=2, max_col=6, values_only=True):
        if any(clean(row[c - 1]) for c in VAL_COLS):
            rules.append(norm(clean(row[NAME_COL - 1])))

names = [l.strip() for l in NAMES.read_text(encoding="utf-8-sig").splitlines() if l.strip()]

missing, hit = Counter(), 0
for n in names:
    k = key_from(n)
    if k in keys or any(has_word(k, r) for r in rules):
        hit += 1
    else:
        missing[k or n] += 1

print(f"الشيت: {SHEET.name}  |  مفاتيح: {len(keys)}  |  قواعد: {len(rules)}")
print(f"الأسماء: {len(names)}  |  معرَّفة: {hit}  |  غير معرَّفة: {len(names) - hit}")
if missing:
    print("\nغير معرَّفة (المفتاح — عدد المهام):")
    for k, c in missing.most_common():
        print(f"  {k}   ({c})")
    OUT.write_text("\ufeff" + "\n".join(f"{k}\t{c}" for k, c in missing.most_common()), encoding="utf-8")
    print(f"\n✅ اتكتب: {OUT.resolve()}")
else:
    print("\n✅ كل المهام معرَّفة — مفيش حاجة ناقصة")