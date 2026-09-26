# -*- coding: utf-8 -*-
"""
build_activity_map.py — يولّد activityMap.js من شيت التصنيفات (تاب واحد).

التخطيط المتوقع في التاب:
  A = اسم المهمة        (زي ما هو في النظام)
  B = تصنيف النشاط   C = نوع النشاط   D = تفاصيل النشاط   E = النوع   F = اسم النوع
  G وما بعده = أعمدة زي «نوع الحدث» → تُتجاهَل تمامًا

السكربت:
  • يستخرج «الكلمة المفتاحية» أوتوماتيك من اسم المهمة
    (يشيل «استمارة مهمة» + الموقع/المحافظة + «مفتوحة اليوم …» + التواريخ).
  • يدمج الصفوف بنفس الكلمة المفتاحية (المهام اليومية المتكررة تبقى صف واحد).
  • يولّد activityMap.js + نسخة من الملف فيها تابين: «الكلمة المفتاحية» و«مراجعة».

التشغيل:
  python build_activity_map.py "C:\\...\\تصنيفات مهام.xlsx"
"""
import sys, re, json
from pathlib import Path
import shutil, tempfile


try:
    from openpyxl import load_workbook
    from openpyxl.styles import Font
except ImportError:
    sys.exit("محتاج openpyxl:  pip install openpyxl")

# ───────────── إعدادات ─────────────
SRC         = Path(sys.argv[1] if len(sys.argv) > 1 else "Book1.xlsx")
OUT         = Path(__file__).resolve().parent / "activityMap.js"   # يُكتب جنب Dashboard.jsx
REVIEW      = SRC.with_name(SRC.stem + "_مراجعة.xlsx")             # نسخة للمراجعة (الأصلي مش بيتغير)
SHEET_NAME  = None            # None = أول تاب | أو اكتب اسم التاب بالنص
NAME_COL    = 1               # A  = اسم المهمة
VAL_COLS    = [2, 3, 4, 5, 6] # B–F = التصنيفات الخمسة
DEDUP_MODE  = "first"         # first = أول صف يكسب | last = آخر صف يكسب
MIN_KEY_LEN = 3               # أقل طول مقبول للكلمة المفتاحية
RULES_SHEET = "قواعد"         # تاب اختياري لكلمات عامة (تأمين، إسعافات…) — لو مش موجود يُتجاهل

# ───────────── تطبيع موحّد (نفس المنطق في activityMap.js بالحرف) ─────────────
DASH   = re.compile(r"[-\u2010-\u2015\u2043\u2212]")
DAYISH = re.compile(r"^(?:(?:ال|و)[\u0600-\u06FF]*|\d+)$")

def norm(s):
    s = str(s if s is not None else "")
    s = re.sub(r"[\u064B-\u0652\u0640\u0670]", "", s)   # تشكيل + تطويل
    s = re.sub(r"[أإآٱ]", "ا", s)
    s = re.sub(r"[ىئ]", "ي", s)
    s = re.sub(r"ة", "ه", s)
    s = re.sub(r"[\"“”«»'‘’]", " ", s)
    s = DASH.sub("-", s)
    return re.sub(r"\s+", " ", s).strip()

def _strip_open_day(s):
    """يشيل «مفتوحة + رقم اليوم» اللي جاية في الآخر من غير شرطة قبلها."""
    i = s.rfind("مفتوحه")
    if i == -1:
        return s
    tail = s[i + 6:].strip()
    toks = tail.split() if tail else []
    if toks and len(toks) <= 6 and all(DAYISH.match(t) for t in toks):
        return s[:i].strip()
    return s

def key_from_norm(s):
    """s مُطبَّع مسبقًا — نفس المنطق بالحرف في الجافاسكربت."""
    for _ in range(3):
        s = re.sub(r"\s*مفتوحه\s*(اليوم.*)?$", "", s)
        s = re.sub(r"\s*\(\s*اليوم[^)]*\)\s*$", "", s)
        s = re.sub(r"\s*\d{4}\s*-\s*\d{1,2}\s*-\s*\d{1,2}\s*$", "", s).strip()
    s = _strip_open_day(s)
    parts = re.split(r"[اأإآ]\s*ستماره\s+مهمه", s)
    if len(parts) > 1:
        s = parts[-1].strip()
    return re.split(r"\s*-\s*", s, maxsplit=1)[0].strip()

def clean(v):
    return re.sub(r"\s+", " ", str(v)).strip() if v is not None else ""

# ───────────── القراءة ─────────────
def load_source(path):
    try:
        return load_workbook(path, data_only=True)
    except PermissionError:
        tmp = Path(tempfile.gettempdir()) / (path.stem + "_tmp.xlsx")
        shutil.copy2(path, tmp)
        print("ℹ️ الملف الأصلي مقفول (مفتوح في Excel؟) — قرأت نسخة مؤقتة.")
        return load_workbook(tmp, data_only=True)

wb = load_source(SRC)
sheet = wb[SHEET_NAME] if SHEET_NAME else wb.worksheets[0]
sheet_title = sheet.title

data, order, counts = {}, [], {}
total = skipped = short = dup_rows = 0
conflicts = {}          # key -> set(tuple(vals))
review_rows = []
max_col = max([NAME_COL] + VAL_COLS)

for row in sheet.iter_rows(min_row=2, max_col=max_col, values_only=True):
    total += 1
    raw = clean(row[NAME_COL - 1])
    vals = [clean(row[c - 1]) for c in VAL_COLS]
    if not raw or not any(vals):
        skipped += 1
        continue

    key = key_from_norm(norm(raw))
    if len(key) < MIN_KEY_LEN:
        short += 1
        review_rows.append([raw, key, "كلمة مفتاحية قصيرة/فارغة — عايزة صف واضح في الشيت"])
        continue

    if key in data:
        dup_rows += 1
        counts[key] = counts.get(key, 1) + 1
        if tuple(vals) != tuple(data[key]):
            conflicts.setdefault(key, set()).add(tuple(vals))
        if DEDUP_MODE == "first":
            continue
    else:
        order.append(key)
        counts[key] = 1
    data[key] = vals

# ───────────── قواعد عامة (اختياري) — تُطبَّق فقط لو مفيش مطابقة كاملة ─────────────
rules = []
if RULES_SHEET and RULES_SHEET in wb.sheetnames:
    for row in wb[RULES_SHEET].iter_rows(min_row=2, max_col=max_col, values_only=True):
        raw = clean(row[NAME_COL - 1])
        vals = [clean(row[c - 1]) for c in VAL_COLS]
        if raw and any(vals):
            rules.append([norm(raw)] + vals)

# ───────────── activityMap.js ─────────────
js_rows = ",\n".join("  " + json.dumps([k] + data[k], ensure_ascii=False) for k in order)

TEMPLATE = r"""// ⚠️ ملف مولَّد آلياً من شيت التصنيفات — متعدّلوش بالإيد.
// المفتاح = الكلمة المفتاحية المستخرجة (بدون «استمارة مهمة» / الموقع / «مفتوحة اليوم …»).
// يرجع: [تصنيف النشاط, نوع النشاط, تفاصيل النشاط, النوع, اسم النوع]

const ROWS = [
__ROWS__
];

// قواعد عامة (اختياري) — تُطبَّق فقط عند فشل المطابقة الكاملة. [الكلمة, 5 تصنيفات]
const RULES = [
__RULES__
];
const norm = (s) => String(s == null ? '' : s)
  .replace(/[\u064B-\u0652\u0640\u0670]/g, '')
  .replace(/[أإآٱ]/g, 'ا')
  .replace(/[ىئ]/g, 'ي')
  .replace(/[ة]/g, 'ه')
  .replace(/["“”«»'‘’]/g, ' ')
  .replace(/[\u2010-\u2015\u2043\u2212]/g, '-')
  .replace(/\s+/g, ' ')
  .trim();

const DAYISH = /^(?:(?:ال|و)[\u0600-\u06FF]*|\d+)$/;

const stripOpenDay = (s) => {
  const i = s.lastIndexOf('مفتوحه');
  if (i === -1) return s;
  const tail = s.slice(i + 6).trim();
  const toks = tail ? tail.split(/\s+/) : [];
  if (toks.length && toks.length <= 6 && toks.every((t) => DAYISH.test(t))) return s.slice(0, i).trim();
  return s;
};

const keyFrom = (rawName) => {
  let s = norm(rawName);
  for (let i = 0; i < 3; i++) {
    s = s.replace(/\s*مفتوحه\s*(اليوم.*)?$/, '')
         .replace(/\s*\(\s*اليوم[^)]*\)\s*$/, '')
         .replace(/\s*\d{4}\s*-\s*\d{1,2}\s*-\s*\d{1,2}\s*$/, '')
         .trim();
  }
  s = stripOpenDay(s);
  const parts = s.split(/[اأإآ]\s*ستماره\s+مهمه/);
  if (parts.length > 1) s = parts[parts.length - 1].trim();
  return s.split(/\s*-\s*/)[0].trim();
};

// فهرس مرة واحدة عند فتح الصفحة + كاش ⇒ التصدير مش بيبطّأ
const MAP = new Map();
for (const r of ROWS) {
  if (r[0] && !MAP.has(r[0])) MAP.set(r[0], r.slice(1));
}
const CACHE = new Map();

const hasWord = (hay, needle) => (' ' + hay + ' ').indexOf(' ' + needle + ' ') !== -1;

export const classifyActivity = (missionName) => {
  let k = CACHE.get(missionName);
  if (k === undefined) { k = keyFrom(missionName); CACHE.set(missionName, k); }
  const hit = MAP.get(k);
  if (hit) return hit.slice();
  for (const r of RULES) {
    if (hasWord(k, r[0])) return r.slice(1);
  }
  return ['', '', '', '', ''];
};
"""

js_rules = ",\n".join("  " + json.dumps(r, ensure_ascii=False) for r in rules)
OUT.write_text(TEMPLATE.replace("__ROWS__", js_rows).replace("__RULES__", js_rules), encoding="utf-8")
# ───────────── تابين المراجعة (في نسخة من الملف) ─────────────
for name in ("الكلمة المفتاحية", "مراجعة"):
    if name in wb.sheetnames:
        del wb[name]

ws1 = wb.create_sheet("الكلمة المفتاحية")
ws1.append(["الكلمة المفتاحية", "تصنيف النشاط", "نوع النشاط", "تفاصيل النشاط", "النوع", "اسم النوع", "عدد الصفوف", "حالة"])
for k in order:
    ws1.append([k] + data[k] + [counts.get(k, 1), "⚠️ تعارض" if k in conflicts else ""])
for c in ws1[1]:
    c.font = Font(bold=True)
ws1.freeze_panes = "A2"

ws2 = wb.create_sheet("مراجعة")
ws2.append(["اسم المهمة الأصلي", "الكلمة المفتاحية", "السبب / التصنيفات المتعارضة"])
for k, variants in conflicts.items():
    for v in sorted(variants):
        ws2.append(["", k, "نفس الكلمة المفتاحية بتصنيف مختلف: " + " | ".join(v)])
for r in review_rows:
    ws2.append(r)
for c in ws2[1]:
    c.font = Font(bold=True)
ws2.freeze_panes = "A2"

try:
    wb.save(REVIEW)
    saved = str(REVIEW)
except PermissionError:
    alt = OUT.parent / (SRC.stem + "_مراجعة.xlsx")
    wb.save(alt)
    saved = str(alt) + "   ⚠️ (المكان الأصلي كان مقفول)"
    wb.close()

# ───────────── التقرير ─────────────
print(f"✅ الملف المولَّد: {OUT}")
print(f"✅ نسخة المراجعة: {saved}")
print(f"   التاب المقروء:                {sheet_title}")
print(f"   صفوف الشيت:                    {total}")
print(f"   صفوف ناقصة (اتخطّت):           {skipped}")
print(f"   صفوف بكلمة مفتاحية قصيرة:      {short}")
print(f"   صفوف اتدمجت (نفس المفتاح):     {dup_rows}")
print(f"   كلمات مفتاحية فريدة:           {len(order)}")
print(f"   ⚠️ مفاتيح بتعارض تصنيف:         {len(conflicts)}")
for k in order[:3]:
    print("   نموذج:", k, "→", " | ".join(data[k]))
